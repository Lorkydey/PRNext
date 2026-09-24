"""Create a local Excel dashboard and Word report from a completed benchmark.

Usage: python scripts/report-next-comparison.py <measurements.json> <output-dir>
Dependencies: scripts/report-requirements.txt. No cloud service is used.
"""
import argparse
import hashlib
import json
import statistics
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import xlsxwriter
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = '#172D43'
NEXT = '#61728A'
RUST = '#008B7C'
LIGHT = '#F0F5F8'
GRID = '#DFE7EE'
AMBER = '#B87413'
NAMES = {
    'public-32k': 'Fichier public · 32 Kio',
    'pages-static': 'Pages · statique',
    'app-static': 'App · statique',
    'pages-ssr': 'Pages · SSR',
    'app-ssr': 'App · SSR',
    'pages-api': 'API Pages',
    'app-api': 'API App',
}
BASE = list(NAMES)
GLOBAL = 'mixed-total'


def fr(value, digits=1):
    return f'{value:,.{digits}f}'.replace(',', '\u202f').replace('.', ',')


def winner(next_value, rust_value, lower=True):
    delta = rust_value / next_value - 1
    if abs(delta) < .05:
        return 'Proche'
    return 'Rustyx' if (delta < 0) == lower else 'Next.js'


def validate(data):
    assert data['status'] == 'complete', 'Benchmark is incomplete'
    assert data['engines'] == ['next', 'rustyx'], 'Report expects two engines'
    assert data['method']['repetitions'] == 3
    assert len(data['runs']) == 78
    assert len(data['builds']) == 6 and len(data['browser']) == 4
    assert all(b['hydrationPassed'] for b in data['browser'])
    for row in data['runs']:
        assert row['errors'] == 0, 'Investigate failed requests before publishing'
        assert not row['reachedRequestCap']
        for key in ['requestsPerSecond', 'loadMedianRssMiB', 'loadPeakRssMiB', 'cpuMsPerRequest']:
            assert row[key] > 0, (row['engine'], row['scenario'], key)
        if row['scenario'] == GLOBAL:
            assert row['requests'] == 70000
            assert len(row['endpointCounts']) == 7
            assert set(row['endpointCounts'].values()) == {10000}


class Report:
    def __init__(self, source, output):
        self.source = source
        self.out = output
        self.data = json.loads(source.read_text())
        validate(self.data)
        self.out.mkdir(parents=True, exist_ok=True)
        self.charts = self.out / 'graphiques'
        self.charts.mkdir(exist_ok=True)
        self.date = datetime.fromisoformat(self.data['finishedAt'].replace('Z', '+00:00')).strftime('%d/%m/%Y')
        self.rows = {(e, s): [r for r in self.data['runs'] if r['engine'] == e and r['scenario'] == s]
                     for e in ['next', 'rustyx'] for s in BASE + [GLOBAL] + [s + '-gzip' for s in BASE[:3]] + ['pages-ssr-c8', 'app-ssr-c8']}

    def med(self, engine, scenario, key):
        return statistics.median(row[key] for row in self.rows[engine, scenario])

    def pair(self, scenario, key, divisor=1):
        return [self.med(e, scenario, key) / divisor for e in ['next', 'rustyx']]

    def gain(self, scenario, key):
        n, r = self.pair(scenario, key)
        return 1 - r / n

    def build(self, engine, key):
        return statistics.median(row[key] for row in self.data['builds'] if row['engine'] == engine)

    def plot_pair(self, name, metric, xlabel, digits=1):
        plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10, 'text.color': NAVY,
                             'axes.labelcolor': NAVY, 'xtick.color': NEXT, 'ytick.color': NAVY})
        fig, ax = plt.subplots(figsize=(9.4, 4.15), dpi=220)
        y = np.arange(len(BASE))
        n = [self.med('next', s, metric) for s in BASE]
        r = [self.med('rustyx', s, metric) for s in BASE]
        for values, offset, color, label in [(n, -.18, NEXT, 'Next.js'), (r, .18, RUST, 'Rustyx')]:
            bars = ax.barh(y + offset, values, height=.30, color=color, label=label)
            ax.bar_label(bars, labels=[fr(v, digits) for v in values], padding=4, fontsize=8.5, color=NAVY)
        ax.set_yticks(y, list(NAMES.values()))
        ax.invert_yaxis()
        ax.set_xlim(0, max(n + r) * 1.19)
        ax.set_xlabel(xlabel, fontsize=10, labelpad=10)
        ax.xaxis.grid(True, color=GRID, linewidth=.6)
        ax.set_axisbelow(True)
        ax.tick_params(axis='both', length=0)
        for spine in ax.spines.values():
            spine.set_visible(False)
        ax.legend(loc='lower left', bbox_to_anchor=(0, 1.01), frameon=False, ncol=2)
        fig.tight_layout(pad=.8)
        target = self.charts / (name + '.png')
        fig.savefig(target, bbox_inches='tight', facecolor='white')
        plt.close(fig)
        return target

    def plots(self):
        self.plot_pair('memoire-par-route', 'loadMedianRssMiB', 'RAM totale sous charge (Mio) — moins = mieux')
        self.plot_pair('cpu-par-route', 'cpuMsPerRequest', 'Temps CPU par réponse valide (ms) — moins = mieux', 3)
        self.plot_pair('debit-par-route', 'requestsPerSecond', 'Réponses valides par seconde — plus = mieux', 0)
        fig, ax = plt.subplots(figsize=(9.4, 2.9), dpi=220)
        metrics = [('Temps pour 70 000 requêtes', 'elapsedMs'), ('CPU total pour cette charge', 'serverCpuMs'), ('RAM globale médiane', 'loadMedianRssMiB')]
        y = np.arange(3)
        r = [self.med('rustyx', GLOBAL, key) / self.med('next', GLOBAL, key) * 100 for _, key in metrics]
        ax.barh(y - .17, [100] * 3, height=.28, color=NEXT, label='Next.js = 100')
        bars = ax.barh(y + .17, r, height=.28, color=RUST, label='Rustyx')
        ax.bar_label(bars, labels=[fr(v, 1) for v in r], padding=5, fontsize=10, color=NAVY)
        ax.set_yticks(y, [label for label, _ in metrics])
        ax.invert_yaxis()
        ax.set_xlim(0, max(100, *r) * 1.2)
        ax.xaxis.grid(True, color=GRID)
        ax.set_axisbelow(True)
        for spine in ax.spines.values():
            spine.set_visible(False)
        ax.tick_params(axis='both', length=0)
        ax.set_xlabel('Indice par critère · Next.js = 100 · moins = mieux')
        ax.legend(loc='lower left', bbox_to_anchor=(0, 1.01), frameon=False, ncol=2, fontsize=9)
        fig.tight_layout(pad=.8)
        fig.savefig(self.charts / 'vue-globale.png', bbox_inches='tight', facecolor='white')
        plt.close(fig)

    def excel(self):
        target = self.out / 'Rustyx-vs-Nextjs.xlsx'
        wb = xlsxwriter.Workbook(target)
        wb.set_properties({'title': 'Rustyx / Next.js — comparaison mesurée', 'author': 'Rustyx',
                           'comments': 'Mesures locales ; données et hypothèses incluses. Aucun score composite CPU/RAM.'})
        f = {
            'title': wb.add_format({'font_name': 'Aptos Display', 'font_size': 22, 'bold': True, 'font_color': 'white', 'bg_color': NAVY, 'valign': 'vcenter'}),
            'sub': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'font_color': NEXT, 'text_wrap': True, 'valign': 'vcenter'}),
            'head': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'bold': True, 'font_color': 'white', 'bg_color': NAVY, 'text_wrap': True, 'valign': 'vcenter'}),
            'text': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'font_color': NAVY, 'bottom': 1, 'bottom_color': GRID, 'valign': 'vcenter'}),
            'wrap': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'font_color': NAVY, 'text_wrap': True, 'valign': 'vcenter', 'bottom': 1, 'bottom_color': GRID}),
            'num': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'num_format': '#,##0.0', 'font_color': NAVY, 'bottom': 1, 'bottom_color': GRID}),
            'precise': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'num_format': '0.000', 'font_color': NAVY, 'bottom': 1, 'bottom_color': GRID}),
            'int': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'num_format': '#,##0', 'font_color': NAVY, 'bottom': 1, 'bottom_color': GRID}),
            'pct': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'num_format': '0.0%;[Red]-0.0%', 'font_color': RUST, 'bottom': 1, 'bottom_color': GRID}),
            'ratio': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'num_format': '0.00"×"', 'font_color': NAVY, 'bottom': 1, 'bottom_color': GRID}),
            'green': wb.add_format({'bg_color': '#E1F3EE', 'font_color': '#006D5F'}),
            'amber': wb.add_format({'bg_color': '#FFF0DC', 'font_color': '#8C5709'}),
            'note': wb.add_format({'font_name': 'Aptos', 'font_size': 10, 'font_color': NEXT, 'text_wrap': True, 'bg_color': LIGHT, 'valign': 'vcenter'}),
        }
        def sheet(name, subtitle, end_col=9):
            ws = wb.add_worksheet(name)
            ws.hide_gridlines(2)
            ws.set_zoom(90)
            ws.set_tab_color(RUST if name in ['Synthèse', 'Global', 'Mémoire', 'CPU'] else NEXT)
            ws.set_column(0, 0, 29)
            ws.set_column(1, end_col, 15)
            ws.set_row(0, 32)
            ws.merge_range(0, 0, 1, end_col, name.upper() + '  /  RUSTYX × NEXT.JS', f['title'])
            ws.merge_range(2, 0, 3, end_col, subtitle, f['sub'])
            ws.freeze_panes(6, 1)
            ws.set_landscape()
            ws.set_paper(9)
            ws.fit_to_pages(1, 0)
            ws.set_margins(.3, .3, .4, .4)
            ws.set_footer('&LComparatif local • ' + self.date + '&RPage &P / &N')
            return ws
        def heads(ws, row, labels):
            ws.write_row(row, 0, labels, f['head'])
            ws.set_row(row, 32)
        def gains(ws, first, last, col):
            ws.conditional_format(first, col, last, col, {'type': 'cell', 'criteria': '>=', 'value': .05, 'format': f['green']})
            ws.conditional_format(first, col, last, col, {'type': 'cell', 'criteria': '<=', 'value': -.05, 'format': f['amber']})
        def chart(ws, row, col, data_sheet, first, last, cols, title, axis, categories_col=0, height=370):
            c = wb.add_chart({'type': 'bar'})
            for index, column in enumerate(cols):
                c.add_series({'name': ['Next.js', 'Rustyx'][index],
                              'categories': [data_sheet, first, categories_col, last, categories_col],
                              'values': [data_sheet, first, column, last, column],
                              'fill': {'color': [NEXT, RUST][index]}, 'border': {'none': True}})
            c.set_title({'name': title, 'name_font': {'name': 'Aptos', 'size': 13, 'color': NAVY}})
            c.set_x_axis({'name': axis, 'min': 0, 'major_gridlines': {'visible': True, 'line': {'color': GRID}}, 'num_font': {'size': 9}})
            c.set_y_axis({'reverse': True, 'num_font': {'size': 10}})
            c.set_legend({'position': 'bottom', 'font': {'size': 10}})
            c.set_chartarea({'border': {'none': True}, 'fill': {'color': 'white'}})
            c.set_plotarea({'border': {'none': True}})
            c.set_size({'width': 840, 'height': height})
            ws.insert_chart(row, col, c)
        summary = sheet('Synthèse', f'{self.date} · Production · 3 passages · Apple M4 / 16 Gio · Next.js {self.data["versions"]["next"]}', 12)
        glob = sheet('Global', '70 000 requêtes par passage : 10 000 sur chacun des sept parcours, entrelacées sur UNE instance. Quatre connexions simultanées.')
        mem = sheet('Mémoire', 'RSS totale : processus Rust et Node, threads RSC inclus. Client de charge exclu. La RAM des scénarios ne s’additionne pas.')
        cpu = sheet('CPU', 'Temps CPU par réponse valide, tous les processus du serveur. Moins = mieux. Un débit élevé ne garantit pas un faible coût CPU.')
        speed = sheet('Débit', 'Réponses valides par seconde et latence p95, quatre requêtes simultanées. Médianes de trois passages.')
        gzip = sheet('Gzip', 'Même protocole, avec compression HTTP. Les fichiers publics sont compressés à la volée ; les pages construites disposent de variantes précompressées.')
        extras = sheet('Build et web', 'Build du projet de test et JavaScript réellement chargé par Chromium. Installation des dépendances et compilation du binaire Rust exclues.')
        raw = sheet('Brut', 'Chaque passage mesuré. Les tableaux de synthèse référencent ces cellules avec MEDIAN ; les résultats calculés sont aussi stockés dans le fichier.', 22)
        method = sheet('Méthode', 'Périmètre, conventions de calcul, définitions et références des données.')
        fields = [
            ('engine', 'Moteur'), ('scenario', 'Scénario'), ('repetition', 'Passage'), ('concurrency', 'Connexions'),
            ('requests', 'Réponses valides'), ('elapsedMs', 'Durée (ms)'), ('serverCpuMs', 'CPU total (ms)'),
            ('cpuMsPerRequest', 'CPU / réponse (ms)'), ('loadMedianRssMiB', 'RSS médiane (Mio)'),
            ('loadPeakRssMiB', 'RSS pic (Mio)'), ('requestsPerSecond', 'Réponses/s'), ('p50Ms', 'p50 (ms)'),
            ('p95Ms', 'p95 (ms)'), ('p99Ms', 'p99 (ms)'), ('errors', 'Erreurs'), ('attempts', 'Tentatives'),
            ('startupMs', 'Démarrage (ms)'), ('coldRequestMs', '1re requête (ms)'), ('clientCpuMs', 'CPU client (ms)'),
            ('meanBodyBytes', 'Octets/réponse'), ('encoding', 'Encodage'), ('warmRss', 'RSS après chauffe'), ('counts', 'Répartition requêtes'),
        ]
        heads(raw, 5, [label for _, label in fields])
        raw_rows = {}
        ordered = sorted(self.data['runs'], key=lambda r: (r['scenario'], r['engine'], r['repetition']))
        for row_index, row in enumerate(ordered, 6):
            raw_rows.setdefault((row['engine'], row['scenario']), []).append(row_index)
            for column, (key, _) in enumerate(fields):
                value = row['afterEqualWarmup']['rssMiB'] if key == 'warmRss' else json.dumps(row['endpointCounts']) if key == 'counts' else row[key]
                raw.write(row_index, column, value, f['num'] if isinstance(value, (float, int)) else f['text'])
        raw.autofilter(5, 0, 5 + len(ordered), len(fields) - 1)
        raw.set_column(1, 1, 22)
        raw.set_column(22, 22, 60)
        from xlsxwriter.utility import xl_rowcol_to_cell
        def formula(e, s, key, divisor=1):
            column = next(i for i, (k, _) in enumerate(fields) if k == key)
            refs = ["'Brut'!" + xl_rowcol_to_cell(row, column, True, True) for row in raw_rows[e, s]]
            return '=MEDIAN(' + ','.join(refs) + ')' + (f'/{divisor}' if divisor != 1 else '')
        def measure(ws, row, col, e, s, key, fmt='num', divisor=1):
            ws.write_formula(row, col, formula(e, s, key, divisor), f[fmt], self.med(e, s, key) / divisor)
        def label(ws, row, name):
            ws.write(row, 0, name, f['text'])
            ws.set_row(row, 25)
        global_metrics = [('Temps total (s)', 'elapsedMs', 1000), ('CPU total (s)', 'serverCpuMs', 1000),
                          ('RAM médiane (Mio)', 'loadMedianRssMiB', 1), ('Pic RAM (Mio)', 'loadPeakRssMiB', 1)]
        heads(glob, 5, ['Critère', 'Next.js', 'Rustyx', 'Gain Rustyx', 'Avantage'])
        for row, (name, key, div) in enumerate(global_metrics, 6):
            label(glob, row, name)
            for col, engine in enumerate(['next', 'rustyx'], 1): measure(glob, row, col, engine, GLOBAL, key, divisor=div)
            n, r = self.pair(GLOBAL, key)
            glob.write_formula(row, 3, f'=1-C{row+1}/B{row+1}', f['pct'], 1-r/n)
            glob.write(row, 4, winner(n, r), f['text'])
        gains(glob, 6, 9, 3)
        heads(glob, 13, ['Moteur', 'Passage', 'Durée (s)', 'CPU total (s)', 'RSS médiane', 'Pic RSS', 'Réponses', 'p95 (ms)', 'Cœurs équiv.'])
        for row, item in enumerate(sorted(self.rows['next', GLOBAL] + self.rows['rustyx', GLOBAL], key=lambda x:(x['engine'], x['repetition'])), 14):
            glob.write_row(row, 0, [item['engine'], item['repetition'], item['elapsedMs']/1000, item['serverCpuMs']/1000,
                                   item['loadMedianRssMiB'], item['loadPeakRssMiB'], item['requests'], item['p95Ms'], item['serverCpuMs']/item['elapsedMs']], f['num'])
        heads(glob, 23, ['Indice indépendant', 'Next.js = 100', 'Rustyx'])
        for row, (name, key, _) in enumerate(global_metrics[:3], 24):
            label(glob, row, name)
            glob.write(row, 1, 100, f['num'])
            source_row = row - 18
            glob.write_formula(row, 2, f'=100*C{source_row+1}/B{source_row+1}', f['num'], 100*self.med('rustyx', GLOBAL, key)/self.med('next', GLOBAL, key))
        glob.merge_range('A29:J31', 'Les 3 passages répètent chacun la même charge. Le CPU total est un temps processeur cumulé, pas un pourcentage. La RAM est celle d’une instance ayant chargé les routes Pages et App ; elle n’est pas la somme de sept serveurs isolés.', f['note'])
        # Per-route tables, with formulas back to raw measurements.
        for ws, key, fmt, unit in [(mem, 'loadMedianRssMiB', 'num', 'Mio'), (cpu, 'cpuMsPerRequest', 'precise', 'ms/réponse')]:
            heads(ws, 5, ['Parcours', f'Next ({unit})', f'Rustyx ({unit})', 'Gain Rustyx', 'Avantage', 'Next min', 'Next max', 'Rustyx min', 'Rustyx max'])
            for row, (s, name) in enumerate(NAMES.items(), 6):
                label(ws, row, name)
                for col, engine in enumerate(['next', 'rustyx'], 1): measure(ws, row, col, engine, s, key, fmt)
                n, r = self.pair(s, key)
                ws.write_formula(row, 3, f'=1-C{row+1}/B{row+1}', f['pct'], 1-r/n)
                ws.write(row, 4, winner(n, r), f['text'])
                ws.write_row(row, 5, [min(x[key] for x in self.rows['next', s]), max(x[key] for x in self.rows['next', s]),
                                      min(x[key] for x in self.rows['rustyx', s]), max(x[key] for x in self.rows['rustyx', s])], f[fmt])
            gains(ws, 6, 12, 3)
            chart(ws, 15, 0, ws.get_name(), 6, 12, [1, 2], 'Coût mémoire par route' if ws == mem else 'Coût CPU par réponse', f'{unit} — moins = mieux')
            ws.print_area(0, 0, 36, 9)
        heads(speed, 5, ['Parcours', 'Next rép./s', 'Rustyx rép./s', 'Ratio débit', 'Next p95 ms', 'Rustyx p95 ms', 'Avantage débit'])
        for row, (s, name) in enumerate(NAMES.items(), 6):
            label(speed, row, name)
            for col, engine in enumerate(['next', 'rustyx'], 1): measure(speed, row, col, engine, s, 'requestsPerSecond', 'int')
            n, r = self.pair(s, 'requestsPerSecond')
            speed.write_formula(row, 3, f'=C{row+1}/B{row+1}', f['ratio'], r/n)
            for col, engine in enumerate(['next', 'rustyx'], 4): measure(speed, row, col, engine, s, 'p95Ms', 'precise')
            speed.write(row, 6, winner(n, r, False), f['text'])
        chart(speed, 15, 0, 'Débit', 6, 12, [1, 2], 'Débit des réponses valides', 'Réponses/s — plus = mieux')
        heads(speed, 36, ['Test à 8 connexions', 'Erreurs Next', 'Erreurs Rustyx', 'p95 Next ms', 'p95 Rustyx ms'])
        for row, s in enumerate(['pages-ssr-c8', 'app-ssr-c8'], 37):
            label(speed, row, s)
            for col, engine in enumerate(['next', 'rustyx'], 1): speed.write(row, col, sum(x['errors'] for x in self.rows[engine, s]), f['int'])
            for col, engine in enumerate(['next', 'rustyx'], 3): measure(speed, row, col, engine, s, 'p95Ms', 'precise')
        heads(gzip, 5, ['Parcours gzip', 'CPU Next ms', 'CPU Rustyx ms', 'Gain CPU', 'RSS Next Mio', 'RSS Rustyx Mio', 'Next rép./s', 'Rustyx rép./s'])
        for row, s in enumerate(BASE[:3], 6):
            name = NAMES[s]; s += '-gzip'
            label(gzip, row, name)
            for col, engine in enumerate(['next', 'rustyx'], 1): measure(gzip, row, col, engine, s, 'cpuMsPerRequest', 'precise')
            gzip.write_formula(row, 3, f'=1-C{row+1}/B{row+1}', f['pct'], self.gain(s, 'cpuMsPerRequest'))
            for col, engine in enumerate(['next', 'rustyx'], 4): measure(gzip, row, col, engine, s, 'loadMedianRssMiB')
            for col, engine in enumerate(['next', 'rustyx'], 6): measure(gzip, row, col, engine, s, 'requestsPerSecond', 'int')
        gains(gzip, 6, 8, 3)
        chart(gzip, 12, 0, 'Gzip', 6, 8, [1, 2], 'CPU avec compression HTTP', 'ms / réponse valide — moins = mieux', height=310)
        heads(extras, 5, ['Critère', 'Next.js', 'Rustyx', 'Gain Rustyx', 'Unité / périmètre'])
        for row, (label_text, key, div, unit) in enumerate([('Build du projet', 'elapsedMs', 1000, 'secondes'), ('Pic RSS du build', 'peakRssMiB', 1, 'Mio, arbre de processus'), ('Sorties de build', 'outputBytes', 1048576, 'Mio, hors dépendances et binaire')], 6):
            label(extras, row, label_text)
            n, r = [self.build(e, key)/div for e in ['next', 'rustyx']]
            extras.write_row(row, 1, [n, r], f['num'])
            extras.write_formula(row, 3, f'=1-C{row+1}/B{row+1}', f['pct'], 1-r/n)
            extras.write(row, 4, unit, f['wrap'])
        for row, endpoint in enumerate(['/pages-static', '/app-static'], 10):
            label(extras, row, 'JS ' + endpoint)
            n, r = [next(b['scriptEncodedBytes'] for b in self.data['browser'] if b['engine']==e and b['endpoint']==endpoint)/1024 for e in ['next','rustyx']]
            extras.write_row(row, 1, [n, r], f['num'])
            extras.write_formula(row, 3, f'=1-C{row+1}/B{row+1}', f['pct'], 1-r/n)
            extras.write(row, 4, 'Kio compressés ; hydratation OK', f['wrap'])
        extras.set_column(4, 4, 38)
        heads(extras, 15, ['Build brut', 'Passage', 'Durée s', 'Pic RSS Mio', 'Sorties Mio'])
        for row, b in enumerate(self.data['builds'], 16):
            extras.write_row(row, 0, [b['engine'], b['repetition'], b['elapsedMs']/1000, b['peakRssMiB'], b['outputBytes']/1048576], f['num'])
        # Dashboard cards and cross-metric winners.
        for first, last, title, key, source_row in [(0, 3, 'GAIN CPU · CHARGE MIXTE', 'serverCpuMs', 7), (4, 7, 'GAIN RAM · CHARGE MIXTE', 'loadMedianRssMiB', 8), (8, 11, 'TEMPS ÉCONOMISÉ · MÊME CHARGE', 'elapsedMs', 6)]:
            summary.merge_range(5, first, 5, last, title, f['head'])
            card = wb.add_format({'font_name': 'Aptos Display', 'font_size': 29, 'bold': True, 'font_color': RUST if self.gain(GLOBAL,key)>=0 else AMBER,
                                  'bg_color': LIGHT, 'num_format': '0.0%', 'valign': 'vcenter', 'align': 'center'})
            summary.merge_range(6, first, 8, last, '', card)
            summary.write_formula(6, first, f"='Global'!D{source_row+1}", card, self.gain(GLOBAL, key))
        summary.merge_range(10, 0, 11, 11, 'Même charge = 70 000 réponses valides par moteur et par passage. Les trois critères restent distincts : aucun « score total » artificiel ne mélange secondes et mémoire.', f['note'])
        chart(summary, 14, 0, 'Global', 24, 26, [1, 2], 'Vue globale — même quantité de travail', 'Indice Next.js = 100 — moins = mieux', height=300)
        heads(summary, 31, ['Parcours', 'Gain RAM', 'Gain CPU', 'Débit Rustyx/Next', 'Avantage RAM', 'Avantage CPU'])
        for row, (s, name) in enumerate(NAMES.items(), 32):
            label(summary, row, name)
            source_row = row - 26
            summary.write_formula(row, 1, f"='Mémoire'!D{source_row+1}", f['pct'], self.gain(s,'loadMedianRssMiB'))
            summary.write_formula(row, 2, f"='CPU'!D{source_row+1}", f['pct'], self.gain(s,'cpuMsPerRequest'))
            summary.write_formula(row, 3, f"='Débit'!D{source_row+1}", f['ratio'], self.med('rustyx',s,'requestsPerSecond')/self.med('next',s,'requestsPerSecond'))
            summary.write(row, 4, winner(*self.pair(s,'loadMedianRssMiB')), f['text'])
            summary.write(row, 5, winner(*self.pair(s,'cpuMsPerRequest')), f['text'])
        gains(summary,32,38,1); gains(summary,32,38,2)
        summary.merge_range(40,0,42,11,'Vert : économie ≥ 5 %. Orange : surcoût ≥ 5 %. « Proche » : écart inférieur à 5 %, repère de lecture et non test statistique. Les chiffres portent sur une petite application, pas sur toute application Next.js.',f['note'])
        summary.print_area(0,0,42,11)
        summary.fit_to_pages(1,1)
        heads(method, 5, ['Élément', 'Détail'])
        details = self.method_items()
        for row, (key, value) in enumerate(details, 6):
            method.write(row,0,key,f['wrap'])
            method.merge_range(row,1,row,9,value,f['wrap'])
            method.set_row(row, 40 if len(value)>130 else 30)
        for ws in [mem,cpu,speed,gzip,extras,glob]: ws.repeat_rows(0,5)
        wb.close()
        return target

    def method_items(self):
        return [
            ('Date et machine', f'{self.date} · {self.data["machine"]["cpu"]} · {self.data["machine"]["logicalCpus"]} cœurs logiques · {self.data["machine"]["ramGiB"]:.0f} Gio · Node {self.data["machine"]["node"]}'),
            ('Versions', f'Next.js {self.data["versions"]["next"]} avec Turbopack par défaut ; Rustyx {self.data["versions"]["rustyx"]} release ; React/react-dom {self.data["versions"]["react"]}.'),
            ('Même projet', 'Sept parcours : fichier public de 32 Kio, Pages/App statiques et SSR, API Pages/App. Pages de 100 lignes avec compteur hydraté ; API JSON de 20 valeurs. Aucun appel distant.'),
            ('Répétitions', 'Trois passages, ordre des moteurs alterné. Serveur neuf par scénario. Un worker Rustyx ; configuration Next par défaut. Aucun changement du framework pendant les mesures.'),
            ('Parcours isolés', '1 000 requêtes de chauffe puis quatre secondes avec quatre connexions keep-alive. Variantes identity et gzip. Les essais SSR à huit connexions durent deux secondes.'),
            ('Charge globale', 'Une seule instance par moteur : chauffe de 1 400 requêtes puis 70 000 demandes entrelacées, exactement 10 000 par parcours, avec quatre connexions. Cette charge est répétée trois fois.'),
            ('CPU', 'Temps CPU serveur et descendants via ps. CPU/réponse = CPU cumulé ÷ réponses valides. Les threads RSC sont inclus dans leur processus. Le CPU du client est exclu.'),
            ('RAM', 'RSS cumulée Rust + Node, échantillonnée toutes les 150 ms environ. Partages de pages potentiellement comptés plusieurs fois ; ce n’est pas une mémoire physique unique. Les pics brefs peuvent être manqués.'),
            ('Agrégation', 'Médiane des trois passages. RAM sous charge : médiane temporelle, puis médiane des passages. Un pic affiché est la médiane des pics observés. Aucune somme de RAM entre scénarios.'),
            ('Gains', 'Économie = 1 − Rustyx/Next pour CPU, RAM et durée. Ratio de débit = Rustyx/Next. Écart < 5 % qualifié de proche, sans signification statistique garantie.'),
            ('Vérification', 'Statut et marqueur vérifiés dans toutes les réponses ; identifiant variable pour le dynamique. Chaque parcours mixte doit recevoir exactement 10 000 requêtes. Hydratation et clic du compteur dans Chromium.'),
            ('Build et web', 'Sorties supprimées avant build ; cache disque OS non vidé. Installation des dépendances et compilation préalable de Rust exclues. Taille du build hors dépendances : pas une taille de déploiement complète.'),
            ('Limites', 'Microbenchmark local court ; client et serveur partagent la machine. Pas de preuve de stabilité longue durée, de consommation électrique ni de capacité maximale. HMR, PPR, images et applications complexes non mesurés.'),
            ('Données', 'mesures-brutes.json — contient chaque passage, les compteurs par route, les versions, les sources de la fixture et les empreintes. Les graphiques Excel restent modifiables.'),
            ('SHA-256 données', hashlib.sha256(self.source.read_bytes()).hexdigest()),
            ('SHA-256 binaire', self.data['binarySha256']),
            ('SHA-256 fixture', self.data['sourceSha256']),
        ]

    def word(self):
        doc = Document()
        section = doc.sections[0]
        section.page_width = Inches(8.27); section.page_height = Inches(11.69)
        section.top_margin = Inches(.67); section.bottom_margin = Inches(.63)
        section.left_margin = section.right_margin = Inches(.68)
        section.header_distance = section.footer_distance = Inches(.27)
        styles = doc.styles
        for name in ['Normal', 'Body Text']:
            styles[name].font.name = 'Arial'
            styles[name].font.size = Pt(10)
            styles[name].font.color.rgb = RGBColor.from_string(NAVY[1:])
            styles[name].paragraph_format.space_after = Pt(6)
            styles[name].paragraph_format.line_spacing = 1.08
        for name, size in [('Title',30),('Heading 1',21),('Heading 2',13)]:
            styles[name].font.name='Arial'
            styles[name].font.size=Pt(size)
            styles[name].font.color.rgb=RGBColor.from_string(NAVY[1:])
            styles[name].paragraph_format.space_before=Pt(8)
            styles[name].paragraph_format.space_after=Pt(7)
        styles['Subtitle'].font.name = 'Arial'
        styles['Subtitle'].font.size = Pt(11)
        styles['Subtitle'].font.color.rgb = RGBColor.from_string(NEXT[1:])
        header=section.header.paragraphs[0]
        header.text='RUSTYX  /  COMPARATIF DE PERFORMANCE'
        header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string(NEXT[1:])
        footer=section.footer.paragraphs[0]
        footer.text=f'Mesures locales · {self.date}                                           RUSTYX × NEXT.JS  |  '
        for run in footer.runs: run.font.size=Pt(8)
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
        core=doc.core_properties
        core.title='Rustyx vs Next.js — CPU, mémoire et performances'
        core.subject='Comparaison expérimentale en production'
        core.author='Rustyx'
        def p(text, style=None): return doc.add_paragraph(text, style)
        def title(number, text, subtitle):
            p(number + '  /  MESURES LOCALES', 'Subtitle')
            doc.add_heading(text,0 if number=='01' else 1)
            p(subtitle)
        def table(headers, rows, widths=None):
            t=doc.add_table(rows=1, cols=len(headers))
            t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
            if widths:
                for c,w in zip(t.columns,widths): c.width=Inches(w)
            for i,h in enumerate(headers): t.rows[0].cells[i].text=h
            for data_row in rows:
                cells=t.add_row().cells
                for i,value in enumerate(data_row): cells[i].text=str(value)
            for ri,row in enumerate(t.rows):
                trpr=row._tr.get_or_add_trPr(); no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
                for ci,cell in enumerate(row.cells):
                    if widths: cell.width=Inches(widths[ci])
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr(); shading=OxmlElement('w:shd')
                    shading.set(qn('w:fill'), NAVY[1:] if ri==0 else ('F0F5F8' if ri%2 else 'FFFFFF')); tcpr.append(shading)
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        item=OxmlElement('w:'+side);item.set(qn('w:w'),'75');item.set(qn('w:type'),'dxa');margins.append(item)
                    tcpr.append(margins)
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1)
                        for run in para.runs:
                            run.font.size=Pt(9)
                            run.bold=ri==0
                            run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else NAVY[1:])
            p('')
            return t
        def image(name, width=6.83): doc.add_picture(str(self.charts/name),width=Inches(width))
        def gain_text(value): return (fr(abs(value)*100)+' % de moins') if value>=0 else (fr(abs(value)*100)+' % de plus')
        title('01','Rustyx × Next.js','Le bilan global, puis les détails CPU et mémoire. Même application, même machine, trois passages en production.')
        p(f'Next.js {self.data["versions"]["next"]} · Rustyx {self.data["versions"]["rustyx"]} · Apple M4 / 16 Gio · {self.date}', 'Subtitle')
        doc.add_heading('Une même charge de 70 000 requêtes',2)
        data=[]
        for label,key,div,unit in [('Temps total','elapsedMs',1000,'s'),('CPU total','serverCpuMs',1000,'s CPU'),('RAM globale médiane','loadMedianRssMiB',1,'Mio'),('Pic RAM observé','loadPeakRssMiB',1,'Mio')]:
            n,r=self.pair(GLOBAL,key,div)
            data.append([label,fr(n)+' '+unit,fr(r)+' '+unit,gain_text(1-r/n)])
        table(['Critère','Next.js','Rustyx','Écart Rustyx'],data,[2.0,1.25,1.25,2.3])
        image('vue-globale.png')
        doc.add_heading('Ce que le résultat permet de conclure',2)
        cpu_gain=self.gain(GLOBAL,'serverCpuMs');ram_gain=self.gain(GLOBAL,'loadMedianRssMiB');time_gain=self.gain(GLOBAL,'elapsedMs')
        p(f'Sur cette charge mixte, Rustyx utilise {gain_text(cpu_gain)} de temps CPU et {gain_text(ram_gain)} de RAM médiane. Il termine le même travail avec {gain_text(time_gain)} de temps écoulé.')
        losing=[NAMES[s] for s in BASE if winner(*self.pair(s,'cpuMsPerRequest'))=='Next.js']
        p('Next.js conserve un avantage CPU sur : '+', '.join(losing)+'. Les tableaux suivants distinguent ces cas des économies de mémoire et des gains de débit.' if losing else 'Les détails suivants distinguent les gains nets des écarts proches ; aucun score arbitraire ne mélange CPU, RAM et vitesse.')
        p('La mémoire globale est mesurée sur UNE instance qui sert tous les parcours. Les fortes économies d’une page statique isolée ne s’appliquent pas automatiquement à une application ayant aussi chargé React et Node.')
        doc.add_page_break()
        title('02','Mémoire : où se trouve le gain ?','RSS totale du serveur, processus Rust et Node compris. Moins = mieux. Chaque parcours ci-dessous utilise un serveur neuf.')
        image('memoire-par-route.png')
        table(['Parcours','Next (Mio)','Rustyx (Mio)','Économie','Avantage'],
              [[name,fr(self.med('next',s,'loadMedianRssMiB')),fr(self.med('rustyx',s,'loadMedianRssMiB')),fr(self.gain(s,'loadMedianRssMiB')*100)+' %',winner(*self.pair(s,'loadMedianRssMiB'))]for s,name in NAMES.items()],
              [2.1,1.0,1.15,1.15,1.4])
        doc.add_heading('Statique et dynamique ne chargent pas les mêmes composants',2)
        p('Les fichiers et pages pré-rendues peuvent être servis entièrement par Rust. Le rendu dynamique continue à utiliser Node et les composants React : son empreinte est donc beaucoup plus élevée. La charge mixte de la première page mesure cet ensemble réellement chargé.')
        p('La RSS est échantillonnée toutes les 150 ms environ. Elle peut compter plusieurs fois des pages partagées ; ce n’est pas une mesure de mémoire physique unique. Les données Excel incluent les pics et les variations des trois passages.')
        doc.add_page_break()
        title('03','CPU : le coût d’une réponse','Temps processeur cumulé du serveur et de ses descendants, divisé par les réponses valides. Moins = mieux.')
        image('cpu-par-route.png')
        table(['Parcours','Next (ms)','Rustyx (ms)','Gain CPU','Avantage'],
              [[name,fr(self.med('next',s,'cpuMsPerRequest'),3),fr(self.med('rustyx',s,'cpuMsPerRequest'),3),fr(self.gain(s,'cpuMsPerRequest')*100)+' %',winner(*self.pair(s,'cpuMsPerRequest'))]for s,name in NAMES.items()],
              [2.1,1.0,1.15,1.15,1.4])
        doc.add_heading('Comment lire ces chiffres',2)
        p('Un gain positif indique une économie de CPU ; un gain négatif indique un surcoût. Le temps CPU additionne le travail des cœurs : il peut dépasser le temps réellement écoulé. Il ne mesure pas la consommation électrique.')
        cores = [statistics.median(r['serverCpuMs'] / r['elapsedMs'] for r in self.rows[e, GLOBAL]) for e in ['next', 'rustyx']]
        p(f'Pendant la charge mixte, Next mobilise en moyenne {fr(cores[0], 2)} cœur équivalent et Rustyx {fr(cores[1], 2)}. Rustyx travaille plus intensément pendant une durée plus courte : son CPU total baisse, mais son occupation moyenne des cœurs augmente.')
        p('« Proche » désigne un écart inférieur à 5 %, simple repère de lecture. Trois passages courts ne permettent pas d’en faire une conclusion statistique. Les minima et maxima figurent dans Excel.')
        doc.add_page_break()
        title('04','Débit, latence et concurrence','Le débit compte uniquement les réponses valides. Les sept parcours de base utilisent quatre connexions simultanées.')
        image('debit-par-route.png')
        table(['Parcours','Next rép./s','Rustyx rép./s','Ratio débit','p95 Next / Rustyx'],
              [[name,fr(self.med('next',s,'requestsPerSecond'),0),fr(self.med('rustyx',s,'requestsPerSecond'),0),'×'+fr(self.med('rustyx',s,'requestsPerSecond')/self.med('next',s,'requestsPerSecond'),2),fr(self.med('next',s,'p95Ms'),2)+' / '+fr(self.med('rustyx',s,'p95Ms'),2)+' ms']for s,name in NAMES.items()],
              [1.9,1.15,1.25,1.0,1.5])
        doc.add_heading('Huit requêtes simultanées : les deux moteurs répondent',2)
        table(['Scénario','Erreurs Next','Erreurs Rustyx','p95 Next / Rustyx'],
              [[s.replace('-c8',''),sum(r['errors'] for r in self.rows['next',s]),sum(r['errors'] for r in self.rows['rustyx',s]),fr(self.med('next',s,'p95Ms'),2)+' / '+fr(self.med('rustyx',s,'p95Ms'),2)+' ms']for s in ['pages-ssr-c8','app-ssr-c8']],
              [1.9,1.4,1.5,2.0])
        p('Ces essais de deux secondes ne démontrent pas une capacité illimitée. Rustyx garde une file bornée : avec un worker, quatre corps peuvent être chargés et 64 requêtes attendre avant lecture de leur corps. Une file pleine renvoie 503 ; une attente expirée renvoie 504.')
        doc.add_page_break()
        title('05','Compression, build et navigateur','Les variantes gzip complètent les mesures sans compression. Le build et le JavaScript navigateur sont mesurés séparément de la charge serveur.')
        doc.add_heading('CPU avec compression HTTP',2)
        table(['Parcours gzip','Next CPU ms','Rustyx CPU ms','Gain CPU','Avantage'],
              [[NAMES[s],fr(self.med('next',s+'-gzip','cpuMsPerRequest'),3),fr(self.med('rustyx',s+'-gzip','cpuMsPerRequest'),3),fr(self.gain(s+'-gzip','cpuMsPerRequest')*100)+' %',winner(*self.pair(s+'-gzip','cpuMsPerRequest'))]for s in BASE[:3]],
              [2.0,1.2,1.3,1.1,1.2])
        p('Les pages construites utilisent leurs variantes précompressées ; le fichier public est compressé à la demande. Le client décompresse les réponses pour vérifier le contenu. Son CPU est exclu des mesures serveur.')
        doc.add_heading('Construire puis lancer le projet',2)
        extra=[]
        for label,key,div,unit in [('Durée du build','elapsedMs',1000,'s'),('Pic RSS du build','peakRssMiB',1,'Mio'),('Sorties de build','outputBytes',1048576,'Mio')]:
            n,r=[self.build(e,key)/div for e in ['next','rustyx']]
            extra.append([label,fr(n,2)+' '+unit,fr(r,2)+' '+unit,gain_text(1-r/n)])
        for label,s,key in [('Serveur HTTP prêt','public-32k','startupMs'),('Premier Pages SSR, HTTP prêt','pages-ssr','coldRequestMs'),('Premier App SSR, HTTP prêt','app-ssr','coldRequestMs')]:
            n,r=self.pair(s,key);extra.append([label,fr(n)+' ms',fr(r)+' ms',winner(n,r)])
        table(['Mesure','Next.js','Rustyx','Lecture'],extra,[2.55,1.05,1.15,2.05])
        p('Rustyx démarre Node à la première requête dynamique. Le serveur HTTP peut donc démarrer vite tout en ayant une première réponse SSR plus lente. Les durées de cette première réponse excluent le démarrage HTTP.')
        doc.add_heading('JavaScript effectivement chargé',2)
        web=[]
        for endpoint in ['/pages-static','/app-static']:
            n,r=[next(b['scriptEncodedBytes']for b in self.data['browser']if b['engine']==e and b['endpoint']==endpoint)/1024 for e in ['next','rustyx']]
            web.append([endpoint,fr(n)+' Kio',fr(r)+' Kio',gain_text(1-r/n)])
        table(['Page','Next.js','Rustyx','Écart'],web,[2.1,1.35,1.35,2.0])
        p('Le compteur a été hydraté puis cliqué dans Chromium, sans erreur JavaScript. La taille du build exclut les dépendances et le binaire Rustyx : elle n’est pas une taille de déploiement complète. Le build exclut également leur installation et la compilation préalable de Rust.')
        doc.add_page_break()
        title('06','Méthode, périmètre et reproductibilité','Ce rapport compare la version actuelle de Rustyx à Next.js sur un projet synthétique précis. Il ne constitue pas une promesse de gain pour toute application.')
        table(['Paramètre','Protocole'],[(k,v)for k,v in self.method_items()[:6]],[1.5,5.3])
        doc.add_heading('Contrôles et limites',2)
        for text in [
            'Toutes les réponses ont été contrôlées. Les routes dynamiques doivent refléter un identifiant variable. Dans le scénario global, chaque moteur traite exactement 10 000 demandes par route, soit 70 000 par passage.',
            'Les tableaux présentent les médianes de trois passages. La RAM sous charge est la médiane des échantillons, puis la médiane des passages. Le pic est la médiane des maxima échantillonnés. Les résultats individuels sont disponibles dans Excel.',
            'Le client et le serveur partagent la même machine. Le client peut limiter les débits les plus élevés. Il n’y a ni réseau distant, ni CDN, ni TLS, ni base applicative. Aucune consommation électrique n’est déduite du temps CPU.',
            'Les sorties de build sont supprimées avant chaque mesure ; le cache disque du système reste chaud. Le benchmark ne couvre pas les images, le HMR, le pré-rendu partiel, toutes les options Next.js ni une application métier lourde.',
            'Les essais sont courts : ils ne remplacent pas un test d’endurance, une recherche de fuite mémoire ou une validation sur la future application importée.',
        ]: p(text)
        doc.add_heading('Fichiers fournis',2)
        p('Rustyx-vs-Nextjs.xlsx : synthèse, charge globale, détails mémoire et CPU, débit, gzip, build et navigateur, mesures brutes et méthode. Les graphiques Excel et les formules restent modifiables.')
        p('mesures-brutes.json : données complètes, versions, sources de l’application et empreintes SHA-256. Le script du dépôt scripts/bench-next-comparison.mjs permet de relancer la mesure.')
        p('Commande : BENCH_GZIP=1 BENCH_MIXED=1 BENCH_OUTPUT=reports/next-vs-rustyx/mesures-brutes.json RUSTYX_NEXT_REFERENCE=/chemin/vers/node_modules/next npm run bench:next')
        doc.save(self.out/'Rapport-Rustyx-vs-Nextjs.docx')

    def finish(self):
        summary={'date':self.date,'source':str(self.source),'binarySha256':self.data['binarySha256'],
                 'global':{key:{'next':self.med('next',GLOBAL,key),'rustyx':self.med('rustyx',GLOBAL,key),'gain':self.gain(GLOBAL,key)} for key in ['serverCpuMs','elapsedMs','loadMedianRssMiB','loadPeakRssMiB']},
                 'totalRequestsVerified':sum(r['requests']for r in self.data['runs']),
                 'errors':sum(r['errors']for r in self.data['runs'])}
        (self.out/'synthese.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False)+'\n')
        (self.out/'LISEZ-MOI.txt').write_text('Rustyx / Next.js — '+self.date+'\n\nRapport-Rustyx-vs-Nextjs.docx : rapport de lecture.\nRustyx-vs-Nextjs.xlsx : tableaux, formules et graphiques modifiables.\nmesures-brutes.json : chaque mesure et les paramètres du test.\nsynthese.json : résultats globaux extraits.\ngraphiques/ : images des graphiques du rapport.\n\nLe total correspond à 70 000 requêtes identiques sur une instance de chaque moteur. Les RAM des scénarios ne sont pas additionnées. Voir la méthode et les limites dans le rapport.\n')
        print(json.dumps(summary,ensure_ascii=False,indent=2))


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('source',type=Path)
    parser.add_argument('output',type=Path)
    args=parser.parse_args()
    report=Report(args.source,args.output)
    report.plots()
    report.excel()
    report.word()
    report.finish()


if __name__=='__main__': main()
